"""ATS Resume Checker - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-2.5-flash"
MAX_CHARS = 20000  # keep prompt size reasonable

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.
Evaluate the resume text provided. If a job description is given, judge keyword and skill match against it;
otherwise judge it as a general-purpose, ATS-friendly resume.

Return ONLY valid JSON with exactly this structure:
{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {
    "formatting_and_structure": <integer 0-100>,
    "keywords_and_skills": <integer 0-100>,
    "work_experience_impact": <integer 0-100>,
    "education_and_certifications": <integer 0-100>,
    "readability_and_clarity": <integer 0-100>
  },
  "strengths": ["<string>", ...],
  "missing_keywords": ["<string>", ...],
  "improvements": [
    {"priority": "High|Medium|Low", "issue": "<string>", "suggestion": "<string>"}
  ]
}
Be honest and specific. Do not inflate scores. Give 4-8 improvements, most important first."""


# ---------- File reading ----------
def extract_text(uploaded_file) -> str:
    """Extract text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    return text.strip()


# ---------- AI ----------
def get_api_key() -> str | None:
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass  # no secrets file locally
    return os.environ.get("GEMINI_API_KEY")


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences or extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(result: dict) -> dict:
    """Make sure every field exists and has the right type so the UI never crashes."""
    section_labels = {
        "formatting_and_structure": "Formatting & Structure",
        "keywords_and_skills": "Keywords & Skills",
        "work_experience_impact": "Work Experience Impact",
        "education_and_certifications": "Education & Certifications",
        "readability_and_clarity": "Readability & Clarity",
    }
    sections = result.get("section_scores") or {}
    improvements = []
    for item in result.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).capitalize(),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
        elif isinstance(item, str):
            improvements.append({"priority": "Medium", "issue": "", "suggestion": item})

    return {
        "ats_score": _clamp(result.get("ats_score")),
        "summary": str(result.get("summary", "")),
        "section_scores": {
            label: _clamp(sections.get(key)) for key, label in section_labels.items()
        },
        "strengths": [str(s) for s in (result.get("strengths") or [])],
        "missing_keywords": [str(k) for k in (result.get("missing_keywords") or [])],
        "improvements": improvements,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    client = genai.Client(api_key=api_key)
    prompt = f"RESUME:\n{resume_text[:MAX_CHARS]}\n\n"
    if job_description.strip():
        prompt += f"JOB DESCRIPTION:\n{job_description[:MAX_CHARS]}\n"
    else:
        prompt += "JOB DESCRIPTION: (none provided - evaluate generally)\n"

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            response_mime_type="application/json",
            temperature=0.2,
        ),
    )
    return normalize_result(parse_json_response(response.text))


# ---------- UI ----------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good - room to improve"
    if score >= 40:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["ats_score"]
    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score}/100")
        st.progress(score / 100)
        st.caption(score_label(score))
    with col2:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.subheader("Score Breakdown")
    for label, value in result["section_scores"].items():
        st.write(f"**{label}** - {value}/100")
        st.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("Missing Keywords")
        for k in result["missing_keywords"] or ["None identified."]:
            st.markdown(f"- {k}")

    st.subheader("Suggested Improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    if not result["improvements"]:
        st.info("No improvements returned.")
    for imp in result["improvements"]:
        icon = icons.get(imp["priority"], "🟠")
        st.markdown(f"{icon} **{imp['priority']} priority** - {imp['issue']}")
        st.markdown(f"> {imp['suggestion']}")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_api_key()
    if not api_key:
        st.error(
            "Gemini API key not found. Set `GEMINI_API_KEY` in `.streamlit/secrets.toml` "
            "(or in Streamlit Cloud secrets) or as an environment variable."
        )
        st.stop()

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, for a tailored score)", height=150
    )

    if st.button("Analyze Resume", type="primary", disabled=uploaded is None):
        try:
            text = extract_text(uploaded)
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return
        if len(text) < 50:
            st.error(
                "Could not extract enough text. If your PDF is a scanned image, "
                "upload a text-based PDF or DOCX instead."
            )
            return
        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(text, job_description, api_key)
            except Exception as e:
                st.error(f"Analysis failed: {e}")
                return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
